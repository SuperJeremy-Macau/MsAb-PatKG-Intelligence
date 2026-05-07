from __future__ import annotations

import json
import os
import re
import shutil
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in os.sys.path:
    os.sys.path.insert(0, str(ROOT))

from bsab_kg_qa_en.config import load_settings
from bsab_kg_qa_en.core.llm_provider import LLMProvider
from bsab_kg_qa_en.intents import IntentRegistry
from bsab_kg_qa_en.interaction import QueryClarifier
from bsab_kg_qa_en.kg import Neo4jRunner
from bsab_kg_qa_en.ner import NERService


SETTINGS_PATH = ROOT / "bsab_kg_qa_en" / "config" / "settings.yaml"
LOCAL_SETTINGS_TXT = ROOT / "bsab_kg_qa_en" / "config" / "Local Settings.txt"
SOURCE_XLSX = ROOT / "dataset6_manual_edit_20260414.xlsx"
OUTPUT_XLSX = ROOT / "dataset6_manual_edit_20260414_entity_surface_test.xlsx"
OUTPUT_SHEET = "entity_surface_test"


@dataclass
class TestCase:
    entity_type: str
    source_row_id: str
    source_intent: str
    expected_slot: str
    expected_value: str
    natural_question: str
    source_question: str


TEST_CASES: list[TestCase] = [
    TestCase(
        entity_type="assignee",
        source_row_id="custom-template",
        source_intent="TARGETPAIRS_BY_ASSIGNEE",
        expected_slot="assignee",
        expected_value="ZHONGSHAN AKESO BIOPHARMA CO LTD",
        natural_question="For the company Akeso, which target-pair combinations are covered in the patent graph?",
        source_question="Custom minimal assignee-input template for interactive extraction.",
    ),
    TestCase(
        entity_type="target",
        source_row_id="custom-template",
        source_intent="ASSIGNEES_BY_TARGET",
        expected_slot="target",
        expected_value="PD-1",
        natural_question="Which assignees are working on PD1 in the current patent graph?",
        source_question="Custom minimal target-input template for interactive extraction.",
    ),
    TestCase(
        entity_type="target_pair",
        source_row_id="custom-template",
        source_intent="ASSIGNEES_BY_TARGETPAIR",
        expected_slot="tp_name",
        expected_value="PD-1/VEGFA",
        natural_question="Which assignees filed patents related to the PD1 and VEGF target-pair combination?",
        source_question="Custom target-pair-input template, aligned with ASSIGNEES_BY_TARGETPAIR phrasing.",
    ),
    TestCase(
        entity_type="functional_of_target",
        source_row_id="S1-006",
        source_intent="TOP_TARGETPAIRS_BY_FUNCTION_PATENT_COUNT",
        expected_slot="functional_of_target",
        expected_value="Adaptive_Immune_Checkpoint_Target",
        natural_question="Within target-pair combinations involving adaptive immune checkpoint targets, which pairs rank highest by patent count? Return the top 10.",
        source_question="Within target-pair combinations involving Adaptive_Immune_Checkpoint_Target, which pairs rank highest by patent count? Return the top 10.",
    ),
    TestCase(
        entity_type="pathway",
        source_row_id="custom-template",
        source_intent="TARGETPAIRS_BY_PATHWAY",
        expected_slot="pathway",
        expected_value="Drug ADME",
        natural_question="Which target-pair combinations include at least one target from the drug absorption, distribution, metabolism, and excretion pathway?",
        source_question="Custom pathway-input template using a natural-language variant of Drug ADME.",
    ),
    TestCase(
        entity_type="technologyclass1",
        source_row_id="TC-S1-005",
        source_intent="TOP_TARGETPAIRS_BY_TECHCLASS1_PATENT_COUNT",
        expected_slot="technologyclass1",
        expected_value="Extending PK/PD",
        natural_question="Using the current patent graph, within the extending PK and PD technology class, which target-pair combinations rank highest by patent count? Return the top 10.",
        source_question="Within the Extending PK/PD TechnologyClass1 category, which target-pair combinations rank highest by patent count? Return the top 10.",
    ),
    TestCase(
        entity_type="origin",
        source_row_id="S4-004",
        source_intent="TOP_ORIGIN_BY_CAGR_5Y",
        expected_slot="origin",
        expected_value="US",
        natural_question="Please answer from the current dataset: within the United States origin group, which organizations or individuals show the strongest five-year filing CAGR?",
        source_question="Please answer from the current dataset: within the US origin group, which organizations or individuals show the strongest five-year filing CAGR?",
    ),
    TestCase(
        entity_type="cancer",
        source_row_id="S1-077",
        source_intent="HIGH_EXPRESSION_TARGETS_BY_CANCER",
        expected_slot="cancer",
        expected_value="BRCA",
        natural_question="Which targets show high expression in breast cancer?",
        source_question="Which targets show high expression in ACC?",
    ),
]


def load_api_key_from_local_settings() -> None:
    if os.getenv("OPENAI_API_KEY"):
        return
    if not LOCAL_SETTINGS_TXT.exists():
        return
    text = LOCAL_SETTINGS_TXT.read_text(encoding="utf-8", errors="ignore")
    match = re.search(r"OPENAI_API_KEY\s*[:=]\s*[\"']?([^\"'\r\n]+)", text)
    if match:
        os.environ["OPENAI_API_KEY"] = match.group(1).strip()


def build_clarifier() -> QueryClarifier:
    cfg = load_settings(str(SETTINGS_PATH))
    neo = cfg["neo4j"]
    llm_cfg = cfg["llm"]
    intent_cfg = cfg["intent"]

    runner = Neo4jRunner(
        uri=neo["uri"],
        user=neo["user"],
        password=neo["password"],
        database=neo["database"],
        max_rows=int(neo.get("max_rows", 50)),
    )
    registry = IntentRegistry(intent_cfg["definitions_dir"])
    llm = LLMProvider(
        base_url=llm_cfg["base_url"],
        api_key_env=llm_cfg["api_key_env"],
        model=llm_cfg["model"],
    )
    ner_service = NERService(runner, llm)
    clarifier = QueryClarifier(runner=runner, registry=registry, ner_service=ner_service, llm=llm)
    return clarifier


def top_candidates(slot: Any, limit: int = 5) -> list[str]:
    if not slot:
        return []
    values = []
    for item in slot.candidates[:limit]:
        label = item.label or item.value
        source = item.source or ""
        values.append(f"{label} [{source}]")
    return values


def evaluate_case(clarifier: QueryClarifier, case: TestCase) -> dict[str, Any]:
    draft = clarifier.start(case.natural_question)
    slot = draft.slots.get(case.expected_slot)
    candidates = top_candidates(slot)
    selected_value = slot.selected_value if slot else None
    candidate_values = [item.value for item in slot.candidates] if slot else []
    passed = selected_value == case.expected_value or case.expected_value in candidate_values
    return {
        "entity_type": case.entity_type,
        "source_row_id": case.source_row_id,
        "source_intent": case.source_intent,
        "expected_slot": case.expected_slot,
        "expected_value": case.expected_value,
        "natural_question": case.natural_question,
        "source_question": case.source_question,
        "flow_state": draft.flow_state,
        "selected_intent": draft.selected_intent,
        "slot_present": bool(slot),
        "slot_status": slot.status if slot else None,
        "raw_text": slot.raw_text if slot else None,
        "selected_value": selected_value,
        "top_candidates": candidates,
        "pass": passed,
        "error": draft.error_message,
    }


def write_results(results: list[dict[str, Any]]) -> None:
    shutil.copyfile(SOURCE_XLSX, OUTPUT_XLSX)
    wb = load_workbook(OUTPUT_XLSX)
    if OUTPUT_SHEET in wb.sheetnames:
        del wb[OUTPUT_SHEET]
    ws = wb.create_sheet(OUTPUT_SHEET)
    headers = [
        "entity_type",
        "source_row_id",
        "source_intent",
        "expected_slot",
        "expected_value",
        "natural_question",
        "source_question",
        "flow_state",
        "selected_intent",
        "slot_present",
        "slot_status",
        "raw_text",
        "selected_value",
        "top_candidates",
        "pass",
        "error",
    ]
    for col, header in enumerate(headers, start=1):
        ws.cell(1, col, header)
    for row_idx, result in enumerate(results, start=2):
        row = [
            result["entity_type"],
            result["source_row_id"],
            result["source_intent"],
            result["expected_slot"],
            result["expected_value"],
            result["natural_question"],
            result["source_question"],
            result["flow_state"],
            result["selected_intent"],
            result["slot_present"],
            result["slot_status"],
            result["raw_text"],
            result["selected_value"],
            " | ".join(result["top_candidates"]),
            result["pass"],
            result["error"],
        ]
        for col, value in enumerate(row, start=1):
            ws.cell(row_idx, col, value)
    wb.save(OUTPUT_XLSX)


def main() -> None:
    load_api_key_from_local_settings()
    clarifier = build_clarifier()
    try:
        results = [evaluate_case(clarifier, case) for case in TEST_CASES]
    finally:
        clarifier.runner.close()
    write_results(results)
    summary = {
        "output_xlsx": str(OUTPUT_XLSX),
        "sheet": OUTPUT_SHEET,
        "passed": sum(1 for item in results if item["pass"]),
        "total": len(results),
        "results": results,
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
